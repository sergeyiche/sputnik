"""Format-specific text extractors for knowledge base sources."""

from __future__ import annotations

from pathlib import Path


class ConversionError(Exception):
    """Raised when a source file cannot be converted."""


def extract_txt(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ConversionError(f"Cannot decode text file: {path}")


def extract_md(path: Path) -> str:
    return extract_txt(path)


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ConversionError("python-docx is not installed") from exc

    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ConversionError("pypdf is not installed") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- Page {index} ---\n{text}")
    if not pages:
        raise ConversionError(f"No extractable text in PDF: {path}")
    return "\n\n".join(pages)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def extract_xlsx(path: Path) -> str:
    """Extract spreadsheet text as tab-separated rows (all sheets)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ConversionError("openpyxl is not installed") from exc

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — surface as ConversionError
        raise ConversionError(f"Cannot open XLSX: {path}: {exc}") from exc

    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"--- Sheet: {sheet.title} ---")
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_to_str(cell) for cell in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not any(cells):
                    continue
                parts.append("\t".join(cells))
                row_count += 1
            if row_count == 0:
                parts.append("(пусто)")
            parts.append("")
    finally:
        workbook.close()

    text = "\n".join(parts).strip()
    # Only sheet titles / empty markers → treat as empty
    body_lines = [
        line
        for line in text.splitlines()
        if line.strip() and not line.startswith("--- Sheet:") and line.strip() != "(пусто)"
    ]
    if not body_lines:
        raise ConversionError(f"No extractable text in XLSX: {path}")
    return text


EXTRACTORS = {
    ".txt": extract_txt,
    ".md": extract_md,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".xlsx": extract_xlsx,
}

SUPPORTED_EXTENSIONS = frozenset(EXTRACTORS)
