"""Unit tests for knowledge source converters."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from packages.etl.convert import convert_paths
from packages.etl.converters import ConversionError, extract_xlsx


def test_extract_xlsx_reads_rows(tmp_path: Path) -> None:
    path = tmp_path / "doctors.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Врачи"
    ws.append(["ФИО", "Город", "Телефон"])
    ws.append(["Иванов И.И.", "Москва", "+7 900 000-00-00"])
    ws.append(["Петров П.П.", "СПб", None])
    wb.save(path)

    text = extract_xlsx(path)
    assert "--- Sheet: Врачи ---" in text
    assert "ФИО" in text
    assert "Иванов И.И." in text
    assert "Москва" in text
    assert "+7 900 000-00-00" in text


def test_extract_xlsx_empty_raises(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    wb.save(path)
    with pytest.raises(ConversionError, match="No extractable text"):
        extract_xlsx(path)


def test_convert_paths_single_file(tmp_path: Path) -> None:
    sources = tmp_path / "sources"
    sources.mkdir()
    output = tmp_path / "out"
    xlsx = sources / "contacts.xlsx"

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["name", "phone"])
    ws.append(["Clinic", "123"])
    wb.save(xlsx)

    results = convert_paths([xlsx], output, sources_dir=sources, force=True)
    assert len(results) == 1
    assert results[0].status == "created"
    assert results[0].output.exists()
    body = results[0].output.read_text(encoding="utf-8")
    assert "source-format: xlsx" in body
    assert "Clinic" in body
    assert "123" in body
